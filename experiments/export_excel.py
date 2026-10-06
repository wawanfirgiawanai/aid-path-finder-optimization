"""Export every experiment result into ONE Excel workbook plus journal-quality figures.

Called automatically at the end of ``run_experiments.py --phase report`` (so the VS Code task and the
in-app "Run paper experiments" button both produce it). It can also be run on its own:

    python experiments/export_excel.py                      # experiments/results -> OptiRoute_Results.xlsx
    python experiments/export_excel.py --results <dir>/results --figures <dir>/figures

Output
    <results>/OptiRoute_Results.xlsx      one workbook, one sheet per experiment
    <figures>/journal/fig_j*.{png,pdf,svg} 600-dpi raster + vector figures (Times, IEEE column widths)
"""
import argparse
import json
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
DATASET = PROJECT / 'dataset'

ALGOS = ['dijkstra', 'astar', 'bfs', 'ga', 'aco', 'hybrid']
NAMES = {'dijkstra': 'Dijkstra', 'astar': 'A*', 'bfs': 'BFS', 'ga': 'GA', 'aco': 'ACO', 'hybrid': 'Hybrid A*-GA'}
MAIN_GROUPS = ['SA-SD', 'SA-MD', 'MA-SD', 'MA-MD']
SCALE_GROUPS = ['SC-05', 'SC-10', 'SC-20', 'SC-30']
GROUP_LABEL = {'SA-SD': 'Single agent, single destination', 'SA-MD': 'Single agent, multiple destinations',
               'MA-SD': 'Multiple agents, single destination', 'MA-MD': 'Multiple agents, multiple destinations',
               'SC-05': 'Scalability, 5 destinations', 'SC-10': 'Scalability, 10 destinations',
               'SC-20': 'Scalability, 20 destinations', 'SC-30': 'Scalability, 30 destinations'}

# Categorical palette in fixed order (validated for colour-vision deficiency on adjacent pairs);
# every series also has its own marker so identity never relies on colour alone.
COLOR = {'dijkstra': '#2a78d6', 'astar': '#eb6834', 'bfs': '#1baf7a', 'ga': '#eda100', 'aco': '#e87ba4',
         'hybrid': '#008300'}
MARKER = {'dijkstra': 'o', 'astar': 's', 'bfs': '^', 'ga': 'D', 'aco': 'v', 'hybrid': '*'}
XL_MARKER = {'dijkstra': 'circle', 'astar': 'square', 'bfs': 'triangle', 'ga': 'diamond', 'aco': 'x',
             'hybrid': 'star'}
INK, INK2, GRID, NEUTRAL, BAD = '#1a1a1a', '#52514e', '#dedede', '#b4b2a9', '#d03b3b'

FONT = 'Arial'
HEAD_FILL, HEAD_INK, BAND_FILL, HYB_FILL = '1F3864', 'FFFFFF', 'F3F5F9', 'E3F1E3'

# column label, number format (percent columns are converted to fractions before writing)
COLS = {
    'case_id': ('Test case', None), 'group': ('Scenario', None), 'scenario': ('Scenario', None),
    'n_dest': ('Destinations', '0'), 'algo': ('Algo id', None), 'algorithm': ('Algorithm', None),
    'runs': ('Runs', '0'), 'cases': ('Cases', '0'),
    'dist_best': ('Best distance (m)', '#,##0.0'), 'dist_mean': ('Mean distance (m)', '#,##0.0'),
    'dist_std': ('SD distance (m)', '#,##0.0'), 'risk_mean': ('Mean risk exposure', '0.0000'),
    'cost_best': ('Best cost', '#,##0.00'), 'cost_mean': ('Mean cost', '#,##0.00'),
    'cost_std': ('SD cost (runs)', '#,##0.00'), 'cost_std_runs': ('Mean SD cost (runs)', '#,##0.00'),
    'time_mean_ms': ('Mean time (ms)', '#,##0.00'), 'time_std_ms': ('SD time (ms)', '#,##0.00'),
    'optimum': ('Proven optimum', '#,##0.00'), 'gap_pct': ('Gap to optimum', '0.00%'),
    'hit_opt': ('Runs at optimum', '0'), 'vs_dijkstra_pct': ('Cost vs Dijkstra', '+0.00%;-0.00%;0.00%'),
    'best_vs_dijkstra_pct': ('Best case vs Dijkstra', '+0.00%;-0.00%;0.00%'),
    'level1_ms': ('Level 1 A* matrix (ms)', '#,##0.00'), 'level2_ms': ('Level 2 GA (ms)', '#,##0.00'),
    'subset': ('Subset', None), 'vs': ('Hybrid vs', None), 'hybrid_better': ('Hybrid better', '0'),
    'ties': ('Ties', '0'), 'hybrid_worse': ('Hybrid worse', '0'),
    'mean_change_pct': ('Mean cost change', '+0.00%;-0.00%;0.00%'),
    'largest_saving_pct': ('Largest saving', '+0.00%;-0.00%;0.00%'),
    'p': ('p (Wilcoxon)', '0.00E+00'), 'p_holm': ('p (Holm)', '0.00E+00'), 'A12': ('A12 effect size', '0.000'),
    'seed': ('Seed', '0'), 'config_id': ('Config', '0'), 'phase': ('Phase', None),
    'cost': ('Cost', '#,##0.00'), 'distance': ('Distance (m)', '#,##0.00'),
    'risk_exposure': ('Risk exposure', '0.0000'), 'max_risk': ('Max edge risk', '0.00'),
    'high_risk_distance': ('High-risk distance (m)', '#,##0.0'), 'makespan_cost': ('Makespan cost', '#,##0.00'),
    'time_ms': ('Time (ms)', '#,##0.00'), 'agents_used': ('Agents used', '0'), 'covered': ('Service visits (older runs)', '0'),
    'destinations_covered': ('Unique destinations covered', '0'), 'service_visits': ('Service visits', '0'),
    'matrix_ms': ('Matrix time (ms)', '#,##0.00'), 'allocation_ms': ('Allocation time (ms)', '#,##0.00'),
    'routing_ms': ('Routing time (ms)', '#,##0.00'), 'fallback_used': ('Fallback used', '@'),
    'fallback_type': ('Fallback type', '@'), 'fallback_count': ('Fallback count', '0'),
    'generations': ('Generations', '0'), 'mean_norm_cost': ('Mean normalised cost', '0.00000'),
    'mean_time_ms': ('Mean time (ms)', '#,##0.0'), 'budget': ('Budget (evaluations)', '#,##0'),
    'selected': ('Selected', None),
}
PCT_COLS = {'gap_pct', 'vs_dijkstra_pct', 'best_vs_dijkstra_pct', 'mean_change_pct', 'largest_saving_pct'}


# ----------------------------------------------------------------------------------------------------------- figures
def _style():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        'font.family': 'serif', 'font.serif': ['Times New Roman', 'STIXGeneral', 'DejaVu Serif'],
        'mathtext.fontset': 'stix', 'font.size': 8, 'axes.titlesize': 8, 'axes.labelsize': 8,
        'xtick.labelsize': 7, 'ytick.labelsize': 7, 'legend.fontsize': 7, 'axes.linewidth': 0.6,
        'axes.edgecolor': INK2, 'axes.labelcolor': INK, 'xtick.color': INK2, 'ytick.color': INK2,
        'xtick.major.width': 0.6, 'ytick.major.width': 0.6, 'xtick.major.size': 2.5, 'ytick.major.size': 2.5,
        'axes.spines.top': False, 'axes.spines.right': False, 'axes.grid': True, 'grid.color': GRID,
        'grid.linewidth': 0.5, 'axes.axisbelow': True, 'legend.frameon': False, 'savefig.dpi': 600,
        'savefig.bbox': 'tight', 'savefig.pad_inches': 0.02, 'pdf.fonttype': 42, 'svg.fonttype': 'none'})
    return plt


def _save(fig, out, name):
    paths = []
    for ext in ('png', 'pdf', 'svg'):
        p = out / f'{name}.{ext}'
        fig.savefig(p)
        paths.append(p)
    return paths[0]


def make_figures(tcases, tscale, stats, scale_runs, fig_dir):
    """Six publication figures (IEEE single column 3.5 in, double column 7.16 in)."""
    plt = _style()
    from matplotlib.lines import Line2D
    fig_dir.mkdir(parents=True, exist_ok=True)
    made = []

    # Fig. J1: per-case cost relative to Dijkstra, small multiples per main scenario
    others = ['astar', 'bfs', 'ga', 'aco', 'hybrid']
    fig, axes = plt.subplots(1, 4, figsize=(7.16, 2.15), sharey=True)
    rng = np.random.default_rng(7)
    for ax, g in zip(axes, MAIN_GROUPS):
        for i, a in enumerate(others):
            v = tcases[(tcases.group == g) & (tcases.algo == a)].vs_dijkstra_pct.to_numpy()
            y = len(others) - 1 - i
            ax.scatter(v, y + rng.uniform(-0.16, 0.16, len(v)), s=40 if a == 'hybrid' else 14, marker=MARKER[a], color=COLOR[a],
                       edgecolors='white', linewidths=0.4, zorder=3)
            ax.plot([v.mean()] * 2, [y - 0.32, y + 0.32], color=INK, lw=1.1, zorder=4)
        ax.axvline(0, color=INK2, lw=0.6, ls=(0, (3, 2)), zorder=2)
        ax.set_title(f'{g}', pad=3)
        ax.grid(axis='y', visible=False)
        ax.set_xlabel('Cost vs. Dijkstra (%)')
    axes[0].set_yticks(range(len(others)))
    axes[0].set_yticklabels([NAMES[a] for a in reversed(others)])
    axes[0].set_ylim(-0.6, len(others) - 0.4)
    h = [Line2D([], [], ls='none', marker='o', color=NEUTRAL, ms=4, label='Test case (mean of 30 runs)'),
         Line2D([], [], color=INK, lw=1.1, label='Scenario mean'),
         Line2D([], [], color=INK2, lw=0.6, ls=(0, (3, 2)), label='Dijkstra (reference)')]
    fig.legend(handles=h, ncol=3, loc='upper center', bbox_to_anchor=(0.5, 1.07))
    fig.tight_layout(w_pad=0.6)
    made.append(('Fig. J1', 'Cost of each test case relative to Dijkstra in the four main scenarios (negative is '
                 'better). Points: test cases; bars: scenario mean.', _save(fig, fig_dir, 'fig_j1_main_relative')))
    plt.close(fig)

    # Fig. J2: scalability, quality (a) and runtime (b): two panels, never a dual axis
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(7.16, 2.4))
    for a in ALGOS:
        d = tcases[tcases.algo == a]
        d = d[d.group.isin(SCALE_GROUPS)].groupby('n_dest').vs_dijkstra_pct.agg(['mean', 'std', 'count'])
        ci = 2.776 * d['std'] / np.sqrt(d['count'])  # t(0.975, 4) for 5 test cases
        kw = dict(color=COLOR[a], marker=MARKER[a], ms=7 if a == 'hybrid' else 4.2, lw=1.6 if a == 'hybrid' else 1.0,
                  mec='white', mew=0.5, label=NAMES[a], zorder=5 if a == 'hybrid' else 3)
        if a == 'astar':
            kw.update(ls=(0, (4, 2)), mfc='white', mec=COLOR[a])
        a1.plot(d.index, d['mean'], **kw)
        if a == 'hybrid':
            a1.fill_between(d.index, d['mean'] - ci, d['mean'] + ci, color=COLOR[a], alpha=0.12, lw=0)
        t = tscale[tscale.algo == a]
        a2.plot(t.n_dest, t.time_mean_ms, **kw)
    a1.axhline(0, color=INK2, lw=0.6)
    a1.set_ylabel('Cost vs. Dijkstra (%)')
    a2.set_ylabel('Mean computation time (ms)')
    a2.set_yscale('log')
    for ax, tag in ((a1, '(a)'), (a2, '(b)')):
        ax.set_xlabel('Number of destinations (4 agents)')
        ax.set_xticks([5, 10, 20, 30])
        ax.text(0.0, 1.03, tag, transform=ax.transAxes, fontweight='bold', va='bottom')
    h, lab = a1.get_legend_handles_labels()
    fig.legend(h, lab, ncol=6, loc='upper center', bbox_to_anchor=(0.5, 1.08))
    fig.tight_layout(w_pad=1.5)
    made.append(('Fig. J2', 'Scalability on 20 test cases: (a) mean cost relative to Dijkstra (band: 95% confidence '
                 'interval of the hybrid over 5 test cases); (b) mean computation time (log scale).',
                 _save(fig, fig_dir, 'fig_j2_scalability')))
    plt.close(fig)

    # Fig. J3: optimality gap heat map (cases with a proven optimum), single-hue sequential scale
    from matplotlib.colors import LinearSegmentedColormap
    opt = tcases[tcases.optimum.notna()]
    piv = opt.pivot_table(index='case_id', columns='algo', values='gap_pct')[ALGOS]
    piv = piv.loc[sorted(piv.index, key=lambda c: (MAIN_GROUPS.index(c[:5]) if c[:5] in MAIN_GROUPS else 9, c))]
    cmap = LinearSegmentedColormap.from_list('seq', ['#f7fbff', '#c6dbef', '#6baed6', '#2171b5', '#08306b'])
    vmax = max(5.0, float(np.nanpercentile(piv.to_numpy(), 95)))
    fig, ax = plt.subplots(figsize=(3.5, 0.2 * len(piv) + 0.9))
    im = ax.imshow(piv.clip(lower=0).to_numpy(), cmap=cmap, vmin=0, vmax=vmax, aspect='auto')
    for (r, c), v in np.ndenumerate(piv.to_numpy()):
        v = max(0.0, v)
        ax.text(c, r, '0' if v < 0.05 else f'{v:.1f}', ha='center', va='center', fontsize=5.6,
                color='white' if v > vmax * 0.55 else INK)
    ax.set_xticks(range(len(ALGOS)))
    ax.set_xticklabels([NAMES[a].replace('Hybrid A*-GA', 'Hybrid') for a in ALGOS])
    ax.xaxis.tick_top()
    ax.set_yticks(range(len(piv)))
    ax.set_yticklabels(piv.index)
    ax.grid(False)
    for s in ax.spines.values():
        s.set_visible(False)
    ax.tick_params(length=0)
    cb = fig.colorbar(im, ax=ax, fraction=0.05, pad=0.03, extend='max')
    cb.set_label('Mean gap to proven optimum (%)')
    cb.outline.set_linewidth(0.4)
    cb.ax.tick_params(labelsize=6.5)
    fig.tight_layout()
    made.append(('Fig. J3', f'Mean optimality gap (%) of each algorithm on the {len(piv)} test cases whose optimum '
                 'was proven by enumeration (colour scale capped at the 95th percentile).',
                 _save(fig, fig_dir, 'fig_j3_optimality_gap')))
    plt.close(fig)

    # Fig. J4: hybrid saving per test case
    piv = tcases.pivot_table(index=['group', 'case_id'], columns='algo', values='cost_mean')
    sav = (100 * (piv['dijkstra'] - piv['hybrid']) / piv['dijkstra']).reset_index(name='saving')
    order = MAIN_GROUPS + SCALE_GROUPS
    sav['k'] = sav.group.map(order.index)
    sav = sav.sort_values(['k', 'case_id']).reset_index(drop=True)
    fig, ax = plt.subplots(figsize=(7.16, 2.2))
    x = np.arange(len(sav))
    pos = sav.saving > 1e-9
    ax.vlines(x[pos], 0, sav.saving[pos], color=COLOR['hybrid'], lw=1.6)
    ax.scatter(x[pos], sav.saving[pos], s=14, color=COLOR['hybrid'], zorder=3, edgecolors='white', linewidths=0.4)
    ax.scatter(x[~pos], sav.saving[~pos], s=10, facecolors='white', edgecolors=NEUTRAL, linewidths=0.8, zorder=3,
               label='Tie (allocation already optimal)')
    for k, g in enumerate(order):
        idx = sav.index[sav.group == g]
        if k:
            ax.axvline(min(idx) - 0.5, color=GRID, lw=0.8)
        ax.text(float(np.mean(idx)), 1.0, g, transform=ax.get_xaxis_transform(), ha='center', va='bottom', fontsize=7)
    top = sav.saving.max()
    ax.annotate(f'{top:.1f}% ({sav.case_id[sav.saving.idxmax()]})', (sav.saving.idxmax(), top),
                xytext=(6, -2), textcoords='offset points', fontsize=6.5, color=INK, va='top')
    ax.set_xticks(x)
    ax.set_xticklabels([c[-2:] for c in sav.case_id], fontsize=5.5)
    ax.set_xlim(-0.7, len(sav) - 0.3)
    ax.set_ylim(-1, top * 1.12)
    ax.set_xlabel('Test case number within each scenario')
    ax.set_ylabel('Cost saving vs. Dijkstra (%)')
    ax.grid(axis='x', visible=False)
    ax.legend(loc='upper right', bbox_to_anchor=(1.0, 0.93))
    fig.tight_layout()
    made.append(('Fig. J4', 'Cost saving of the hybrid A*-GA over Dijkstra (greedy allocation) on all 40 test cases. '
                 'Hollow points: no saving possible because the allocation was already optimal.',
                 _save(fig, fig_dir, 'fig_j4_hybrid_saving')))
    plt.close(fig)

    # Fig. J5: win / tie / loss of the hybrid against each baseline (Wilcoxon-Holm in labels)
    st = stats[stats.subset == 'all 40 cases'].drop_duplicates('vs')
    fig, ax = plt.subplots(figsize=(3.5, 1.75))
    y = np.arange(len(st))[::-1]
    left = np.zeros(len(st))
    for col, colr, lab in (('hybrid_better', COLOR['hybrid'], 'Hybrid better'), ('ties', NEUTRAL, 'Tie'),
                           ('hybrid_worse', BAD, 'Hybrid worse')):
        v = st[col].to_numpy(dtype=float)
        ax.barh(y, v, left=left, color=colr, height=0.62, edgecolor='white', linewidth=1.0, label=lab)
        for yy, l, w in zip(y, left, v):
            if w >= 3:
                ax.text(l + w / 2, yy, f'{int(w)}', ha='center', va='center', fontsize=6.5,
                        color='white' if colr != NEUTRAL else INK)
        left += v
    for yy, p in zip(y, st.p_holm):
        ax.text(41, yy, f'p = {p:.1e}', va='center', fontsize=6.5, color=INK2)
    ax.set_yticks(y)
    ax.set_yticklabels([f'vs. {v}' for v in st.vs])
    ax.set_xlim(0, 49)
    ax.set_xticks([0, 10, 20, 30, 40])
    ax.set_xlabel('Number of test cases (of 40)')
    ax.grid(axis='y', visible=False)
    ax.legend(ncol=3, loc='upper center', bbox_to_anchor=(0.45, 1.22))
    fig.tight_layout()
    made.append(('Fig. J5', 'Outcome of the hybrid A*-GA against each baseline over 40 test cases, with '
                 'Holm-corrected Wilcoxon signed-rank p-values.', _save(fig, fig_dir, 'fig_j5_win_tie_loss')))
    plt.close(fig)

    # Fig. J6: convergence of the level-2 GA on the largest cases
    conv = scale_runs[(scale_runs.algo == 'hybrid') & (scale_runs.group == 'SC-30')]
    conv = conv.dropna(subset=['convergence']) if 'convergence' in conv else conv.iloc[0:0]
    if len(conv):
        curves = [np.array(json.loads(c), dtype=float) for c in conv.convergence]
        curves = [100 * c / c[0] for c in curves if len(c)]
        n = max(len(c) for c in curves)
        mat = np.vstack([np.pad(c, (0, n - len(c)), mode='edge') for c in curves])
        fig, ax = plt.subplots(figsize=(3.5, 2.1))
        for c in curves:
            ax.plot(range(len(c)), c, color=NEUTRAL, lw=0.4, alpha=0.6)
        q1, med, q3 = np.percentile(mat, [25, 50, 75], axis=0)
        ax.fill_between(range(n), q1, q3, color=COLOR['hybrid'], alpha=0.18, lw=0, label='Interquartile range')
        ax.plot(range(n), med, color=COLOR['hybrid'], lw=1.6, label='Median')
        ax.plot([], [], color=NEUTRAL, lw=0.6, label=f'Single run (n = {len(curves)})')
        ax.set_xlabel('Generation')
        ax.set_ylabel('Best fitness (% of greedy start)')
        ax.set_xlim(0, n - 1)
        ax.legend(ncol=3, loc='lower center', bbox_to_anchor=(0.5, 1.0), handlelength=1.6, columnspacing=1.0)
        fig.tight_layout()
        made.append(('Fig. J6', 'Convergence of the level-2 GA on the five 30-destination test cases '
                     '(fitness normalised by the greedy warm start).', _save(fig, fig_dir, 'fig_j6_convergence')))
        plt.close(fig)
    return made


# ----------------------------------------------------------------------------------------------------------- workbook
def _prep(df):
    df = df.copy()
    for c in df.columns:
        if c in PCT_COLS:
            df[c] = df[c] / 100.0
    return df


class Book:
    def __init__(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.F, self.A, self.P, self.B, self.S = Font, Alignment, PatternFill, Border, Side
        self.thin = Side(style='thin', color='BFBFBF')

    def font(self, **kw):
        return self.F(name=FONT, size=kw.pop('size', 10), **kw)

    def sheet(self, title, heading, note=None, tab=None):
        ws = self.wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        if tab:
            ws.sheet_properties.tabColor = tab
        ws['A1'] = heading
        ws['A1'].font = self.font(size=14, bold=True, color=HEAD_FILL)
        if note:
            ws['A2'] = note
            ws['A2'].font = self.font(size=9, italic=True, color='595959')
        return ws

    def table(self, ws, df, row, col=1, name=None, hybrid_col=None, widths=None, title=None):
        """Write a formatted table; returns (first_data_row, last_data_row)."""
        from openpyxl.utils import get_column_letter
        if title:
            ws.cell(row, col, title).font = self.font(size=11, bold=True, color=HEAD_FILL)
            row += 1
        border = self.B(bottom=self.thin)
        for j, c in enumerate(df.columns):
            cell = ws.cell(row, col + j, COLS.get(c, (c, None))[0])
            cell.font = self.font(bold=True, color=HEAD_INK, size=9.5)
            cell.fill = self.P('solid', fgColor=HEAD_FILL)
            cell.alignment = self.A(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[row].height = 30
        first = row + 1
        hyb_rows = []
        for i, rec in enumerate(df.itertuples(index=False), start=first):
            is_h = hybrid_col is not None and str(getattr(rec, hybrid_col)) in ('hybrid', NAMES['hybrid'])
            for j, (c, v) in enumerate(zip(df.columns, rec)):
                if isinstance(v, float) and np.isnan(v):
                    v = None
                elif isinstance(v, (np.integer,)):
                    v = int(v)
                elif isinstance(v, (np.floating,)):
                    v = float(v)
                elif isinstance(v, (np.bool_, bool)):
                    v = 'yes' if v else ''
                cell = ws.cell(i, col + j, v)
                fmt = COLS.get(c, (c, None))[1]
                if fmt:
                    cell.number_format = fmt
                cell.font = self.font(size=9.5, bold=is_h)
                cell.border = border
                if is_h:
                    cell.fill = self.P('solid', fgColor=HYB_FILL)
                elif (i - first) % 2:
                    cell.fill = self.P('solid', fgColor=BAND_FILL)
                if fmt is None and j > 0:
                    cell.alignment = self.A(horizontal='center')
            if is_h:
                hyb_rows.append(i)
        last = first + len(df) - 1
        for j, c in enumerate(df.columns):
            letter = get_column_letter(col + j)
            w = (widths or {}).get(c) or max(10, min(28, len(COLS.get(c, (c,))[0]) * 0.55 + 4,
                                                     max([len(str(x)) for x in df[c].head(50)] + [0]) + 3))
            if ws.column_dimensions[letter].width is None or ws.column_dimensions[letter].width < w:
                ws.column_dimensions[letter].width = w
        return first, last

    def save(self, path):
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.save(path)


def _xl_chart_common(ch, title, ytitle, xtitle):
    # axis names go in the title (openpyxl axis titles overlap the tick labels in Excel)
    ch.title = title
    ch.x_axis.tickLblPos = 'low'
    try:
        ch.roundedCorners = False
    except AttributeError:
        pass
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = 'b'
    ch.height, ch.width = 8.5, 17
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E0E0E0', w=6350))


def _figures_sheet(bk, made_figs, XLImage):
    """Sheet with the embedded journal figures (left out with --no-figures)."""
    ws = bk.sheet('Figures', 'Journal-quality figures', 'Embedded at screen size; the files in figures/journal are '
                  '600-dpi PNG plus vector PDF and SVG (Times New Roman, IEEE column widths 3.5 in / 7.16 in).',
                  tab='7030A0')
    r = 4
    for tag, cap, png in made_figs:
        ws.cell(r, 1, f'{tag}. {cap}').font = bk.font(size=10, bold=True)
        ws.cell(r + 1, 1, f'File: {png.parent.name}/{png.stem}.png | .pdf | .svg').font = bk.font(size=8.5,
                                                                                                   color='595959')
        img = XLImage(str(png))
        scale = 760 / img.width if img.width > 1 else 1
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        if img.width / img.height < 2.4 and img.width > 500:  # single-column figures at a readable size
            k = 460 / img.width
            img.width, img.height = int(img.width * k), int(img.height * k)
        ws.add_image(img, f'A{r + 3}')
        r += 3 + int(img.height / 20) + 3
    ws.column_dimensions['A'].width = 12


def build_workbook(results, figures, out_path, made_figs):
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: F401  (kept for users extending the book)

    tm_runs = pd.read_csv(results / 'main_runs.csv')
    ts_runs = pd.read_csv(results / 'scale_runs.csv')
    tcases = pd.read_csv(results / 'table_cases.csv')
    tmain = pd.read_csv(results / 'table_main.csv')
    tscale = pd.read_csv(results / 'table_scale.csv')
    gain = pd.read_csv(results / 'table_hybrid_gain.csv')
    stats = pd.read_csv(results / 'stats_cases.csv')
    perc = pd.read_csv(results / 'stats_runs_per_case.csv')
    summary = json.loads((results / 'summary.json').read_text())
    tuning = pd.read_csv(results / 'tuning.csv') if (results / 'tuning.csv').exists() else pd.DataFrame()
    tuned = json.loads((results / 'tuned_params.json').read_text()) if (results / 'tuned_params.json').exists() else {}
    tc_dir = DATASET / 'test_cases'
    catalog = pd.read_csv(tc_dir / 'test_cases_summary.csv')
    points = pd.read_csv(tc_dir / 'test_cases_points.csv')

    bk = Book()
    from openpyxl.utils import get_column_letter
    from openpyxl.comments import Comment

    # ---------------------------------------------------------------- 1. Summary
    ws = bk.sheet('Summary', 'OptiRoute: risk-aware routing experiments, results workbook',
                  f'Generated {datetime.now():%Y-%m-%d %H:%M} from {results}. All numbers are recomputed from the '
                  'raw runs in the "Runs" sheets; simulated scenarios on the OpenStreetMap road network of Mamuju '
                  '(r = 1,500 m).', tab=HEAD_FILL)
    ws.column_dimensions['A'].width = 52
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 64
    r = 4
    ws.cell(r, 1, 'Key results').font = bk.font(size=11, bold=True, color=HEAD_FILL)
    n_gain = len(gain) + 1
    kpis = [
        ('Test cases evaluated', "=COUNTA('Hybrid Gain'!A5:A%d)" % (n_gain + 3), '0', '20 main + 20 scalability'),
        ('Total algorithm runs', "=COUNTA('Runs Main'!A5:A%d)+COUNTA('Runs Scale'!A5:A%d)"
         % (len(tm_runs) + 4, len(ts_runs) + 4), '#,##0', '6 algorithms; 30 runs (main) and 10 runs (scalability)'),
        ('Cases where hybrid beats Dijkstra/A*', "=SUMPRODUCT(ISNUMBER('Hybrid Gain'!G5:G{0})*('Hybrid Gain'!G5:G{0}>0.000001))".format(n_gain + 3), '0',
         'remaining cases are ties: greedy allocation was already optimal'),
        ('Cases where hybrid is worse than Dijkstra/A*', "=SUMPRODUCT(ISNUMBER('Hybrid Gain'!G5:G{0})*('Hybrid Gain'!G5:G{0}<-0.000001))"
         .format(n_gain + 3), '0', ''),
        ('Mean cost saving vs Dijkstra (40 cases)', "=AVERAGE('Hybrid Gain'!G5:G%d)" % (n_gain + 3), '0.00%', ''),
        ('Largest cost saving vs Dijkstra', "=MAX('Hybrid Gain'!G5:G%d)" % (n_gain + 3), '0.00%', ''),
        ('Cases with a proven optimum', "=COUNT('Hybrid Gain'!F5:F%d)" % (n_gain + 3), '0',
         'optimum by exhaustive enumeration (small cases)'),
        ('... of which hybrid hit the optimum in every run', summary.get('cases_hybrid_reached_optimum'), '0',
         'from summary.json'),
        ('Friedman chi-square (6 algorithms, 40 cases)', summary['friedman_all_cases']['statistic'], '0.00', ''),
        ('Friedman p-value', summary['friedman_all_cases']['p'], '0.00E+00', ''),
    ]
    for k, (lab, val, fmt, note) in enumerate(kpis, start=r + 1):
        ws.cell(k, 1, lab).font = bk.font()
        c = ws.cell(k, 2, val)
        c.number_format = fmt
        c.font = bk.font(bold=True, color='000000' if isinstance(val, str) else '0000FF')
        c.alignment = bk.A(horizontal='right')
        ws.cell(k, 3, note).font = bk.font(size=9, color='595959')
        for j in (1, 2, 3):
            ws.cell(k, j).border = bk.B(bottom=bk.thin)
    r = k + 2
    ws.cell(r, 1, 'Sheets in this workbook').font = bk.font(size=11, bold=True, color=HEAD_FILL)
    guide = [('Test Cases', 'catalogue of the 44 test cases (40 evaluated + 4 tuning) and their points'),
             ('Main Scenarios', 'Table III: SA-SD, SA-MD, MA-SD, MA-MD (5 cases each, 30 runs per algorithm)'),
             ('Per Case', 'every test case x algorithm (mean, best, SD, gap to optimum, time)'),
             ('Hybrid Gain', 'hybrid A*-GA vs Dijkstra per test case (saving computed by Excel formulas)'),
             ('Statistics', 'Wilcoxon signed-rank + Holm, Friedman, Vargha-Delaney A12'),
             ('Scalability', '5, 10, 20, 30 destinations with 4 agents (5 cases each, 10 runs)'),
             ('Tuning', 'random-search tuning of GA, ACO and hybrid; selected parameters'),
             ('Figures', 'journal-quality figures (600 dpi; PDF/SVG vector copies in figures/journal)'),
             ('Runs Main / Runs Scale', 'raw output of every single run (for re-analysis)'),
             ('Notes', 'definitions of every metric and the colour/marker legend')]
    if not made_figs:
        guide = [g for g in guide if g[0] != 'Figures']
    for k, (s, d) in enumerate(guide, start=r + 1):
        c = ws.cell(k, 1, s)
        c.hyperlink = f"#'{s.split(' / ')[0]}'!A1"
        c.font = bk.font(color='0563C1', underline='single')
        ws.cell(k, 3, d).font = bk.font(size=9, color='595959')
    r = k + 2
    ws.cell(r, 1, 'Colour key: black = Excel formula, blue = value copied from the run outputs; '
                  'green rows = proposed hybrid A*-GA.').font = bk.font(size=9, italic=True, color='595959')

    # ---------------------------------------------------------------- 2. Test cases
    ws = bk.sheet('Test Cases', 'Test-case catalogue (one road network, different start and destination points)',
                  'Source: dataset/test_cases/test_cases_summary.csv and test_cases_points.csv. '
                  'Each case also has a map in dataset/test_cases/maps/<ID>.png.', tab='5B9BD5')
    cat = catalog.copy()
    cat.insert(2, 'description', cat.group.map(GROUP_LABEL).fillna('Parameter tuning (not evaluated)'))
    COLS.update({'purpose': ('Purpose', None), 'n_depots': ('Depots', '0'), 'agents_per_depot': ('Agents/depot', '0'),
                 'n_agents': ('Agents', '0'), 'n_destinations': ('Destinations', '0'),
                 'force_distribution': ('Force distribution', None), 'point_seed': ('Point seed', '0'),
                 'min_spacing_m': ('Min spacing (m)', '0'), 'description': ('Description', None),
                 'role': ('Role', None), 'label': ('Label', None), 'node_id': ('OSM node id', '0'),
                 'lat': ('Latitude', '0.000000'), 'lng': ('Longitude', '0.000000'), 'agents': ('Agents', '0')})
    f, l = bk.table(ws, cat, 4, widths={'description': 38})
    ws.freeze_panes = 'B5'
    ws.auto_filter.ref = f'A4:{get_column_letter(cat.shape[1])}{l}'
    c0 = cat.shape[1] + 2
    f2, l2 = bk.table(ws, points, 3, col=c0, title='Points (depots and destinations) of every test case')
    ws.auto_filter.ref = f'A4:{get_column_letter(cat.shape[1])}{l}'

    # ---------------------------------------------------------------- 3. Main scenarios
    ws = bk.sheet('Main Scenarios', 'Main scenarios: mean over 5 test cases x 30 runs per algorithm',
                  'Cost = sum of d(1 + r) over the routes; percentages relative to Dijkstra on the same test case '
                  '(negative = cheaper). Gap = distance to the proven optimum.', tab='70AD47')
    tm = _prep(tmain)
    f, l = bk.table(ws, tm, 4, hybrid_col='algorithm')
    ws.freeze_panes = 'C5'
    ws.auto_filter.ref = f'A4:{get_column_letter(tm.shape[1])}{l}'
    # chart data block (pivot) + native, editable Excel chart
    pr = l + 3
    ws.cell(pr - 1, 1, 'Chart data: mean cost vs Dijkstra (%) by scenario').font = bk.font(bold=True, color=HEAD_FILL)
    others = ['astar', 'bfs', 'ga', 'aco', 'hybrid']
    ws.cell(pr, 1, 'Scenario').font = bk.font(bold=True)
    for j, a in enumerate(others, start=2):
        ws.cell(pr, j, NAMES[a]).font = bk.font(bold=True)
    for i, g in enumerate(MAIN_GROUPS, start=pr + 1):
        ws.cell(i, 1, g).font = bk.font()
        for j, a in enumerate(others, start=2):
            v = tm[(tm.scenario == g) & (tm.algorithm == NAMES[a])].vs_dijkstra_pct.iloc[0]
            c = ws.cell(i, j, float(v))
            c.number_format = '0.0%'
            c.font = bk.font(color='0000FF')
    ch = BarChart()
    ch.type = 'col'
    ch.grouping = 'clustered'
    ch.overlap = -10
    ch.gapWidth = 80
    ch.add_data(Reference(ws, min_col=2, max_col=1 + len(others), min_row=pr, max_row=pr + 4), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=pr + 1, max_row=pr + 4))
    for s, a in zip(ch.series, others):
        s.graphicalProperties.solidFill = COLOR[a][1:]
        s.graphicalProperties.line.solidFill = 'FFFFFF'
    _xl_chart_common(ch, 'Mean cost vs Dijkstra by scenario (%, negative = better)', 'Cost vs Dijkstra', 'Scenario')
    ch.y_axis.numFmt = '0%'
    ws.add_chart(ch, f'A{pr + 7}')

    # ---------------------------------------------------------------- 4. Per case
    ws = bk.sheet('Per Case', 'Every test case x algorithm', 'Use the filter buttons (Scenario, Algorithm) to explore. '
                  'Main cases: 30 runs; scalability cases: 10 runs.', tab='70AD47')
    tc = tcases.copy()
    tc['k1'] = tc.group.map((MAIN_GROUPS + SCALE_GROUPS).index)
    tc['k2'] = tc.algo.map(ALGOS.index)
    tc = _prep(tc.sort_values(['k1', 'case_id', 'k2']).drop(columns=['k1', 'k2', 'algo']))
    f, l = bk.table(ws, tc, 4, hybrid_col='algorithm')
    ws.freeze_panes = 'D5'
    ws.auto_filter.ref = f'A4:{get_column_letter(tc.shape[1])}{l}'

    # ---------------------------------------------------------------- 5. Hybrid gain (with formulas)
    ws = bk.sheet('Hybrid Gain', 'Hybrid A*-GA vs Dijkstra on every test case',
                  'Columns D-F come from the runs (blue); G and H are Excel formulas: saving = (Dijkstra - Hybrid) / '
                  'Dijkstra, gap = (Hybrid - optimum) / optimum.', tab='70AD47')
    g2 = gain.copy()
    gmap = tcases.drop_duplicates('case_id').set_index('case_id')[['group', 'n_dest']]
    g2 = g2.join(gmap, on='case_id')
    g2['k'] = g2.group.map((MAIN_GROUPS + SCALE_GROUPS).index)
    g2 = g2.sort_values(['k', 'case_id'])
    hdr = ['Test case', 'Scenario', 'Destinations', 'Dijkstra mean cost', 'Hybrid mean cost', 'Proven optimum',
           'Saving vs Dijkstra', 'Hybrid gap to optimum']
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(4, j, h)
        c.font = bk.font(bold=True, color=HEAD_INK, size=9.5)
        c.fill = bk.P('solid', fgColor=HEAD_FILL)
        c.alignment = bk.A(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[4].height = 30
    for i, rec in enumerate(g2.itertuples(index=False), start=5):
        vals = [rec.case_id, rec.group, int(rec.n_dest), float(rec.dijkstra), float(rec.hybrid),
                None if pd.isna(rec.optimum) else float(rec.optimum)]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(i, j, v)
            c.font = bk.font(size=9.5, color='0000FF' if j >= 4 else '000000')
            if j >= 4:
                c.number_format = '#,##0.00'
        ws.cell(i, 7, f'=IF(D{i}>0,(D{i}-E{i})/D{i},"")').number_format = '0.00%'
        ws.cell(i, 8, f'=IF(F{i}="","",(E{i}-F{i})/F{i})').number_format = '0.00%'
        for j in range(1, 9):
            ws.cell(i, j).border = bk.B(bottom=bk.thin)
            if j in (7, 8):
                ws.cell(i, j).font = bk.font(size=9.5, bold=(j == 7))
            if (i - 5) % 2:
                ws.cell(i, j).fill = bk.P('solid', fgColor=BAND_FILL)
    last = 4 + len(g2)
    from openpyxl.formatting.rule import DataBarRule
    ws.conditional_formatting.add(f'G5:G{last}', DataBarRule(start_type='num', start_value=0, end_type='max',
                                                             color='63BE7B', showValue=True))
    for j, w in enumerate([12, 11, 12, 16, 16, 15, 16, 16], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'B5'
    ws.auto_filter.ref = f'A4:H{last}'
    t = last + 2
    ws.cell(t, 1, 'Mean').font = bk.font(bold=True)
    ws.cell(t, 7, f'=AVERAGE(G5:G{last})').number_format = '0.00%'
    ws.cell(t + 1, 1, 'Cases with saving > 0').font = bk.font(bold=True)
    ws.cell(t + 1, 7, f'=SUMPRODUCT(ISNUMBER(G5:G{last})*(G5:G{last}>0.000001))')
    for rr in (t, t + 1):
        ws.cell(rr, 7).font = bk.font(bold=True)
    ch = BarChart()
    ch.type = 'col'
    ch.gapWidth = 40
    ch.add_data(Reference(ws, min_col=7, min_row=4, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last))
    ch.series[0].graphicalProperties.solidFill = COLOR['hybrid'][1:]
    _xl_chart_common(ch, 'Cost saving of hybrid A*-GA vs Dijkstra per test case (%)', 'Saving', 'Test case')
    ch.y_axis.numFmt = '0%'
    ch.legend = None
    ch.width = 24
    ws.add_chart(ch, 'J4')

    # ---------------------------------------------------------------- 6. Statistics
    ws = bk.sheet('Statistics', 'Statistical tests', 'Wilcoxon signed-rank on per-case mean cost (Holm-corrected '
                  'within each subset); Friedman over all 40 cases; per-case Mann-Whitney U with Vargha-Delaney A12 '
                  'for stochastic baselines (A12 > 0.5 means the hybrid is cheaper).', tab='ED7D31')
    ws.cell(4, 1, 'Friedman test (6 algorithms x 40 test cases)').font = bk.font(size=11, bold=True, color=HEAD_FILL)
    for k, (lab, v, fmt) in enumerate([('Chi-square', summary['friedman_all_cases']['statistic'], '0.000'),
                                       ('p-value', summary['friedman_all_cases']['p'], '0.00E+00'),
                                       ('Test cases', summary['friedman_all_cases']['cases'], '0')], start=5):
        ws.cell(k, 1, lab).font = bk.font()
        c = ws.cell(k, 2, v)
        c.number_format = fmt
        c.font = bk.font(bold=True, color='0000FF')
    f, l = bk.table(ws, _prep(stats), 10, title='Hybrid A*-GA vs each baseline (paired over test cases)')
    ws.auto_filter.ref = f'A11:{get_column_letter(stats.shape[1])}{l}'
    f2, l2 = bk.table(ws, perc, l + 3, title='Per-case comparison against stochastic baselines (runs)')
    ws.column_dimensions['A'].width = 24

    # ---------------------------------------------------------------- 7. Scalability
    ws = bk.sheet('Scalability', 'Scalability: 4 agents, 5-30 destinations (5 test cases x 10 runs per size)',
                  'Level 1 = A* cost matrix between all depots and destinations; level 2 = allocation GA.', tab='ED7D31')
    tsc = _prep(tscale.drop(columns=['algo']))
    f, l = bk.table(ws, tsc, 4, hybrid_col='algorithm')
    ws.freeze_panes = 'D5'
    ws.auto_filter.ref = f'A4:{get_column_letter(tsc.shape[1])}{l}'
    pr = l + 3
    ws.cell(pr - 1, 1, 'Chart data').font = bk.font(bold=True, color=HEAD_FILL)
    ws.cell(pr, 1, 'Destinations').font = bk.font(bold=True)
    ws.cell(pr, 2 + len(ALGOS), 'Destinations').font = bk.font(bold=True)
    for j, a in enumerate(ALGOS):
        ws.cell(pr, 2 + j, f'{NAMES[a]}').font = bk.font(bold=True)
        ws.cell(pr, 3 + len(ALGOS) + j, f'{NAMES[a]}').font = bk.font(bold=True)
    for i, n in enumerate([5, 10, 20, 30], start=pr + 1):
        ws.cell(i, 1, n)
        ws.cell(i, 2 + len(ALGOS), n)
        for j, a in enumerate(ALGOS):
            d = tscale[(tscale.algo == a) & (tscale.n_dest == n)].iloc[0]
            c = ws.cell(i, 2 + j, float(d.vs_dijkstra_pct) / 100)
            c.number_format = '0.0%'
            c.font = bk.font(color='0000FF')
            c = ws.cell(i, 3 + len(ALGOS) + j, float(d.time_mean_ms))
            c.number_format = '#,##0.0'
            c.font = bk.font(color='0000FF')
    ws.cell(pr - 1, 2, 'Cost vs Dijkstra').font = bk.font(italic=True, size=9)
    ws.cell(pr - 1, 3 + len(ALGOS), 'Mean time (ms)').font = bk.font(italic=True, size=9)

    def _line(cols0, title, ytitle, fmt, log=False):
        lc = LineChart()
        lc.add_data(Reference(ws, min_col=cols0, max_col=cols0 + len(ALGOS) - 1, min_row=pr, max_row=pr + 4),
                    titles_from_data=True)
        lc.set_categories(Reference(ws, min_col=1, min_row=pr + 1, max_row=pr + 4))
        for s, a in zip(lc.series, ALGOS):
            s.graphicalProperties.line.solidFill = COLOR[a][1:]
            s.graphicalProperties.line.width = 28575 if a == 'hybrid' else 15875
            if a == 'astar':
                s.graphicalProperties.line.dashStyle = 'dash'
            s.marker.symbol = XL_MARKER[a]
            s.marker.size = 8 if a == 'hybrid' else 6
            s.marker.graphicalProperties.solidFill = COLOR[a][1:]
            s.marker.graphicalProperties.line.solidFill = COLOR[a][1:]
            s.smooth = False
        _xl_chart_common(lc, title, ytitle, 'Number of destinations')
        lc.y_axis.numFmt = fmt
        if log:
            lc.y_axis.scaling.logBase = 10
        return lc

    ws.add_chart(_line(2, 'Cost vs Dijkstra (%) vs number of destinations', 'Cost vs Dijkstra', '0%'), f'A{pr + 7}')
    ws.add_chart(_line(3 + len(ALGOS), 'Mean time (ms, log scale) vs number of destinations', 'Mean time (ms)', '#,##0',
                       log=True), f'H{pr + 7}')

    # ---------------------------------------------------------------- 8. Tuning
    ws = bk.sheet('Tuning', 'Parameter tuning (random search, 24 configurations x 5 runs on separate tuning cases)',
                  'Selected = lowest mean normalised cost; ties broken by the smaller computational budget. '
                  'GA and ACO were tuned on TUNE-A/B, the hybrid on TUNE-C/D.', tab='A5A5A5')
    r = 4
    ws.cell(r, 1, 'Selected parameters (used in the evaluation)').font = bk.font(size=11, bold=True, color=HEAD_FILL)
    r += 1
    for algo, ps in tuned.items():
        for k, v in ps.items():
            ws.cell(r, 1, NAMES.get(algo, algo)).font = bk.font(bold=algo == 'hybrid')
            ws.cell(r, 2, k).font = bk.font()
            c = ws.cell(r, 3, v)
            c.font = bk.font(bold=True, color='0000FF')
            r += 1
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 24
    r += 1
    for algo in ('ga', 'aco', 'hybrid'):
        if tuning.empty:
            break
        d = tuning[tuning.algo == algo].dropna(axis=1, how='all').drop(columns=['algo'])
        d = d.sort_values('mean_norm_cost')
        f, l = bk.table(ws, d, r, title=f'{NAMES[algo]}: all configurations (sorted, best first)')
        for i in range(f, l + 1):
            if str(ws.cell(i, list(d.columns).index('selected') + 1).value) == 'yes':
                for j in range(1, d.shape[1] + 1):
                    ws.cell(i, j).fill = bk.P('solid', fgColor=HYB_FILL)
                    ws.cell(i, j).font = bk.font(size=9.5, bold=True)
        r = l + 3

    # ---------------------------------------------------------------- 9. Figures
    if made_figs:
        _figures_sheet(bk, made_figs, XLImage)

    # ---------------------------------------------------------------- 10-11. Raw runs
    for title, df in (('Runs Main', tm_runs), ('Runs Scale', ts_runs)):
        ws = bk.sheet(title, f'Raw runs: {title.split()[1].lower()} experiment ({len(df):,} rows)',
                      'One row per (test case, algorithm, seed). Convergence histories are omitted here; see '
                      'scale_runs.csv.', tab='A5A5A5')
        d = df.drop(columns=[c for c in ('convergence',) if c in df.columns]).copy()
        d['k'] = d.algo.map(ALGOS.index)
        d = d.sort_values(['case_id', 'k', 'seed']).drop(columns=['k'])
        d['algo'] = d.algo.map(NAMES)
        f, l = bk.table(ws, d, 4)
        ws.freeze_panes = 'C5'
        ws.auto_filter.ref = f'A4:{get_column_letter(d.shape[1])}{l}'

    # ---------------------------------------------------------------- 12. Notes
    ws = bk.sheet('Notes', 'Definitions and legend', tab='A5A5A5')
    notes = [
        ('Cost', 'Sum over all routes of d(1 + r), with d = edge length (m) and r in [0, 1] = simulated edge risk.'),
        ('Risk exposure', 'Length-weighted mean risk of the edges travelled.'),
        ('Cost vs Dijkstra', '(cost - Dijkstra cost) / Dijkstra cost on the same test case; negative = cheaper.'),
        ('Gap to optimum', '(cost - proven optimum) / proven optimum; optimum by exhaustive enumeration of the '
                           'allocation on small cases.'),
        ('Runs at optimum', 'number of runs whose cost equals the proven optimum.'),
        ('Wilcoxon / Holm', 'paired signed-rank test on per-case mean cost; Holm correction within each subset.'),
        ('A12', 'Vargha-Delaney effect size; P(hybrid run cheaper than baseline run). 0.5 = no difference, '
                '>= 0.71 = large.'),
        ('Time', 'wall-clock time per run in ms, including the A* cost matrix (level 1) for the hybrid.'),
        ('Data', 'simulated depots, destinations and risk on the OpenStreetMap road network of Mamuju '
                 '(radius 1,500 m, risk seed 42). Only the road geometry comes from OpenStreetMap.'),
        ('Reproduce', 'python experiments/run_experiments.py --phase all --params-file dataset/tuned_params.json '
                      '(this workbook is rebuilt automatically by the report phase).'),
    ]
    for k, (a, b) in enumerate(notes, start=3):
        ws.cell(k, 1, a).font = bk.font(bold=True)
        ws.cell(k, 2, b).font = bk.font()
        ws.cell(k, 2).alignment = bk.A(wrap_text=True, vertical='top')
        ws.cell(k, 1).alignment = bk.A(vertical='top')
    k += 2
    ws.cell(k, 1, 'Series legend (figures and charts)').font = bk.font(size=11, bold=True, color=HEAD_FILL)
    for i, a in enumerate(ALGOS, start=k + 1):
        ws.cell(i, 1, NAMES[a]).font = bk.font(bold=a == 'hybrid')
        ws.cell(i, 2, f'colour {COLOR[a]}, marker {XL_MARKER[a]}').font = bk.font()
        ws.cell(i, 3).fill = bk.P('solid', fgColor=COLOR[a][1:])
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 110
    ws.column_dimensions['C'].width = 6
    ws['B3'].comment = Comment('Risk values are simulated (seed 42); see dataset/README.md.', 'OptiRoute')

    bk.save(out_path)
    return out_path


def export(results, figures, with_figures=True):
    results, figures = Path(results), Path(figures)
    tcases = pd.read_csv(results / 'table_cases.csv')
    tscale = pd.read_csv(results / 'table_scale.csv')
    stats = pd.read_csv(results / 'stats_cases.csv')
    scale_runs = pd.read_csv(results / 'scale_runs.csv')
    made = make_figures(tcases, tscale, stats, scale_runs, figures / 'journal') if with_figures else []
    out = build_workbook(results, figures, results / 'OptiRoute_Results.xlsx', made)
    print('[excel] written', out, flush=True)
    return out


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--results', default=str(ROOT / 'results'))
    ap.add_argument('--figures', default=str(ROOT / 'figures'))
    ap.add_argument('--no-figures', action='store_true',
                    help='build the workbook without the journal figures and the Figures sheet')
    a = ap.parse_args()
    export(a.results, a.figures, with_figures=not a.no_figures)
